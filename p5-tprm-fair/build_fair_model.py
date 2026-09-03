"""
Third-party risk quantification using FAIR, with a Monte Carlo simulation.

Scenario: a payroll SaaS vendor holds employee PII and bank details for 800 employees.
Question the CFO actually asked: "Is this vendor worth the money, or should we bring it back
in-house?" A red/amber/green vendor score cannot answer that. Dollars can.
"""
import random
import statistics as stats
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

random.seed(42)
N = 100_000

def pert(low, likely, high, lam=4.0):
    """BetaPERT sample - the standard FAIR distribution for expert estimates."""
    if high <= low:
        return low
    mu = (low + lam * likely + high) / (lam + 2)
    if abs(mu - likely) < 1e-9:
        a = b = 3.0
    else:
        a = ((mu - low) * (2 * likely - low - high)) / ((likely - mu) * (high - low))
        b = a * (high - mu) / (mu - low)
    a, b = max(a, 0.1), max(b, 0.1)
    return low + random.betavariate(a, b) * (high - low)


SCENARIOS = {
    "current": dict(
        label="Current state — vendor as contracted today",
        tef=(0.8, 2.0, 5.0),          # threat event frequency per year
        vuln=(0.25, 0.40, 0.60),      # probability a threat event becomes a loss event
        primary=(45_000, 120_000, 400_000),
        slef=(0.30, 0.50, 0.75),      # secondary loss event frequency
        secondary=(200_000, 900_000, 3_500_000),
    ),
    "treated": dict(
        label="With contract and control remediation (Option B)",
        tef=(0.8, 2.0, 5.0),
        vuln=(0.10, 0.18, 0.30),
        primary=(30_000, 80_000, 250_000),
        slef=(0.20, 0.35, 0.55),
        secondary=(150_000, 600_000, 2_200_000),
    ),
    "inhouse": dict(
        label="Bring payroll in-house (Option C)",
        tef=(0.5, 1.2, 3.0),
        vuln=(0.15, 0.25, 0.40),
        primary=(40_000, 100_000, 300_000),
        slef=(0.25, 0.40, 0.60),
        secondary=(180_000, 700_000, 2_600_000),
    ),
}


def simulate(p):
    losses = []
    for _ in range(N):
        lef = pert(*p["tef"]) * pert(*p["vuln"])
        events = int(lef) + (1 if random.random() < (lef - int(lef)) else 0)
        total = 0.0
        for _ in range(events):
            total += pert(*p["primary"])
            if random.random() < pert(*p["slef"]):
                total += pert(*p["secondary"])
        losses.append(total)
    losses.sort()
    def pct(q):
        return losses[min(int(q * N), N - 1)]
    return dict(
        mean=sum(losses) / N,
        median=stats.median(losses),
        p10=pct(0.10), p50=pct(0.50), p75=pct(0.75), p90=pct(0.90), p95=pct(0.95), p99=pct(0.99),
        max=losses[-1],
        prob_zero=sum(1 for x in losses if x == 0) / N,
        prob_over_1m=sum(1 for x in losses if x > 1_000_000) / N,
    )


results = {k: simulate(v) for k, v in SCENARIOS.items()}

# ---------------- Questionnaire scoring ----------------
QUESTIONS = [
    ("Q1", "Independent security certification (ISO 27001 / SOC 2 Type II) current within 12 months", 10, 4, 5),
    ("Q2", "Encryption of PII at rest and in transit, documented and verified", 10, 5, 5),
    ("Q3", "MFA enforced on all administrative and remote access", 9, 3, 5),
    ("Q4", "Documented incident response plan, tested within 12 months", 8, 2, 5),
    ("Q5", "Breach notification commitment of 72 hours or less in the contract", 9, 1, 5),
    ("Q6", "Subprocessor list disclosed, with notification of change", 7, 2, 5),
    ("Q7", "Right to audit or to receive audit reports on request", 8, 1, 5),
    ("Q8", "Data return and certified deletion on termination", 7, 3, 5),
    ("Q9", "Background screening of personnel with access to customer data", 6, 4, 5),
    ("Q10", "Vulnerability management with remediation SLAs by severity", 8, 3, 5),
    ("Q11", "Business continuity plan with tested RTO/RPO for payroll runs", 7, 2, 5),
    ("Q12", "Cyber liability insurance at or above contract value", 6, 4, 5),
    ("Q13", "Access reviews performed at least quarterly", 7, 2, 5),
    ("Q14", "Secure development lifecycle with code review and testing", 6, 3, 5),
    ("Q15", "Logging and monitoring retained 12 months or more", 6, 2, 5),
]

wb = Workbook()

hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
money = '"$"#,##0'


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = hdr_fill, hdr_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


# ---- Sheet 1: Executive summary ----
ws = wb.active
ws.title = "Executive Summary"
rows = [
    ["Third-Party Risk Quantification — Payroll SaaS Vendor", ""],
    ["Prepared by", "Moussa Touré, GRC"],
    ["Date", "September 2026"],
    ["Method", "FAIR (Factor Analysis of Information Risk), BetaPERT inputs, 100,000-iteration Monte Carlo"],
    ["", ""],
    ["The question", "Renew the vendor as-is, remediate, or bring payroll in-house?"],
    ["", ""],
    ["Recommendation", "Option B — renew with contract and control remediation."],
    ["Annualised loss expectancy, current", results["current"]["mean"]],
    ["Annualised loss expectancy, remediated", results["treated"]["mean"]],
    ["Risk reduction per year", results["current"]["mean"] - results["treated"]["mean"]],
    ["Cost of remediation (one-off + annual)", 95_000],
    ["Return on control investment", (results["current"]["mean"] - results["treated"]["mean"]) / 95_000],
    ["", ""],
    ["Why not in-house", "Option C removes vendor dependency but shifts the loss to us with no net "
                         "reduction that justifies an estimated $600K build and 9 months. The vendor "
                         "is not the problem; the contract and three controls are."],
    ["Worst-case exposure (P95, current)", results["current"]["p95"]],
    ["Probability of a loss above $1M in a year (current)", results["current"]["prob_over_1m"]],
    ["Probability of a loss above $1M in a year (remediated)", results["treated"]["prob_over_1m"]],
]
for r in rows:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14, color="1F3864")
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
for r in (9, 10, 11, 12, 16):
    ws.cell(row=r, column=2).number_format = money
ws.cell(row=13, column=2).number_format = '0.0"x"'
for r in (17, 18):
    ws.cell(row=r, column=2).number_format = '0.0%'
ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 86

# ---- Sheet 2: FAIR inputs ----
ws2 = wb.create_sheet("FAIR Inputs")
ws2.append(["Scenario", "Factor", "Minimum", "Most likely", "Maximum", "Basis for the estimate"])
style_header(ws2, 6)
BASIS = {
    "Threat event frequency (per year)": "Vendor disclosed 2 security incidents in 3 years; payroll "
        "platforms are a known target for wire-fraud and PII theft.",
    "Vulnerability (loss event given threat event)": "Derived from questionnaire gaps: no MFA on "
        "admin access (Q3), no tested IR plan (Q4), no 72-hour breach notice (Q5).",
    "Primary loss magnitude": "Incident response, forensics, legal, overtime payroll reprocessing.",
    "Secondary loss event frequency": "Probability that regulators, customers or employees respond.",
    "Secondary loss magnitude": "Notification and credit monitoring for 800 employees, regulatory "
        "response, litigation reserve, employee relations impact.",
}
for key, p in SCENARIOS.items():
    for factor, vals in [("Threat event frequency (per year)", p["tef"]),
                         ("Vulnerability (loss event given threat event)", p["vuln"]),
                         ("Primary loss magnitude", p["primary"]),
                         ("Secondary loss event frequency", p["slef"]),
                         ("Secondary loss magnitude", p["secondary"])]:
        ws2.append([p["label"], factor, vals[0], vals[1], vals[2], BASIS[factor]])
        r = ws2.max_row
        if "magnitude" in factor:
            for c in (3, 4, 5):
                ws2.cell(row=r, column=c).number_format = money
        for c in range(1, 7):
            ws2.cell(row=r, column=c).border = border
            ws2.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (1, 6)))
for i, w in enumerate([38, 42, 14, 14, 14, 72], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

# ---- Sheet 3: Simulation results ----
ws3 = wb.create_sheet("Simulation Results")
ws3.append(["Scenario", "Mean (ALE)", "Median", "P10", "P75", "P90", "P95", "P99",
            "P(no loss)", "P(loss > $1M)"])
style_header(ws3, 10)
for key in ("current", "treated", "inhouse"):
    r = results[key]
    ws3.append([SCENARIOS[key]["label"], r["mean"], r["median"], r["p10"], r["p75"],
                r["p90"], r["p95"], r["p99"], r["prob_zero"], r["prob_over_1m"]])
    row = ws3.max_row
    for c in range(2, 9):
        ws3.cell(row=row, column=c).number_format = money
    for c in (9, 10):
        ws3.cell(row=row, column=c).number_format = '0.0%'
    for c in range(1, 11):
        ws3.cell(row=row, column=c).border = border
ws3.column_dimensions["A"].width = 46
for i in range(2, 11):
    ws3.column_dimensions[get_column_letter(i)].width = 15
ws3.append([])
ws3.append(["Read the distribution, not the average. The mean is what to budget; the P95 is what to "
            "survive. A board that only sees the mean is unprepared for the year that goes wrong."])
ws3.cell(row=ws3.max_row, column=1).font = Font(italic=True)

# ---- Sheet 4: Questionnaire scoring ----
ws4 = wb.create_sheet("Vendor Questionnaire")
ws4.append(["ID", "Control question", "Weight", "Vendor response (0-5)", "Max", "Weighted score",
            "Weighted max", "Gap"])
style_header(ws4, 8)
tot_w = tot_max = 0
for qid, text, weight, score, maxs in QUESTIONS:
    ws4.append([qid, text, weight, score, maxs, weight * score, weight * maxs,
                "GAP" if score <= 2 else ""])
    r = ws4.max_row
    tot_w += weight * score
    tot_max += weight * maxs
    if score <= 2:
        for c in range(1, 9):
            ws4.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FCE4E4")
    for c in range(1, 9):
        ws4.cell(row=r, column=c).border = border
        ws4.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c == 2))
ws4.append(["", "TOTAL", "", "", "", tot_w, tot_max, f"{tot_w/tot_max:.0%}"])
ws4.cell(row=ws4.max_row, column=2).font = Font(bold=True)
ws4.cell(row=ws4.max_row, column=8).font = Font(bold=True)
ws4.append([])
ws4.append(["", "Scoring bands: >=85% accept · 70-84% accept with remediation plan · "
                "50-69% conditional, executive approval required · <50% reject"])
ws4.append(["", "This vendor scores in the 'conditional' band. The score alone does not say whether "
                "to renew — the FAIR model does. The questionnaire finds the gaps; the quantification "
                "prices them."])
for r in (ws4.max_row - 1, ws4.max_row):
    ws4.cell(row=r, column=2).font = Font(italic=True)
for i, w in enumerate([7, 66, 9, 20, 8, 15, 14, 8], start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w

wb.save("TPRM-FAIR-Quantification.xlsx")

print(f"Current ALE   : ${results['current']['mean']:,.0f}  (P95 ${results['current']['p95']:,.0f})")
print(f"Treated ALE   : ${results['treated']['mean']:,.0f}  (P95 ${results['treated']['p95']:,.0f})")
print(f"In-house ALE  : ${results['inhouse']['mean']:,.0f}")
print(f"Questionnaire : {tot_w}/{tot_max} = {tot_w/tot_max:.0%}")
