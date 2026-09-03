"""Builds the ISO/IEC 27001:2022 Statement of Applicability workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# (id, name, theme)
CONTROLS = [
    # A.5 Organizational (37)
    ("A.5.1", "Policies for information security", "Organizational"),
    ("A.5.2", "Information security roles and responsibilities", "Organizational"),
    ("A.5.3", "Segregation of duties", "Organizational"),
    ("A.5.4", "Management responsibilities", "Organizational"),
    ("A.5.5", "Contact with authorities", "Organizational"),
    ("A.5.6", "Contact with special interest groups", "Organizational"),
    ("A.5.7", "Threat intelligence", "Organizational"),
    ("A.5.8", "Information security in project management", "Organizational"),
    ("A.5.9", "Inventory of information and other associated assets", "Organizational"),
    ("A.5.10", "Acceptable use of information and other associated assets", "Organizational"),
    ("A.5.11", "Return of assets", "Organizational"),
    ("A.5.12", "Classification of information", "Organizational"),
    ("A.5.13", "Labelling of information", "Organizational"),
    ("A.5.14", "Information transfer", "Organizational"),
    ("A.5.15", "Access control", "Organizational"),
    ("A.5.16", "Identity management", "Organizational"),
    ("A.5.17", "Authentication information", "Organizational"),
    ("A.5.18", "Access rights", "Organizational"),
    ("A.5.19", "Information security in supplier relationships", "Organizational"),
    ("A.5.20", "Addressing information security within supplier agreements", "Organizational"),
    ("A.5.21", "Managing information security in the ICT supply chain", "Organizational"),
    ("A.5.22", "Monitoring, review and change management of supplier services", "Organizational"),
    ("A.5.23", "Information security for use of cloud services", "Organizational"),
    ("A.5.24", "Information security incident management planning and preparation", "Organizational"),
    ("A.5.25", "Assessment and decision on information security events", "Organizational"),
    ("A.5.26", "Response to information security incidents", "Organizational"),
    ("A.5.27", "Learning from information security incidents", "Organizational"),
    ("A.5.28", "Collection of evidence", "Organizational"),
    ("A.5.29", "Information security during disruption", "Organizational"),
    ("A.5.30", "ICT readiness for business continuity", "Organizational"),
    ("A.5.31", "Legal, statutory, regulatory and contractual requirements", "Organizational"),
    ("A.5.32", "Intellectual property rights", "Organizational"),
    ("A.5.33", "Protection of records", "Organizational"),
    ("A.5.34", "Privacy and protection of PII", "Organizational"),
    ("A.5.35", "Independent review of information security", "Organizational"),
    ("A.5.36", "Compliance with policies, rules and standards for information security", "Organizational"),
    ("A.5.37", "Documented operating procedures", "Organizational"),
    # A.6 People (8)
    ("A.6.1", "Screening", "People"),
    ("A.6.2", "Terms and conditions of employment", "People"),
    ("A.6.3", "Information security awareness, education and training", "People"),
    ("A.6.4", "Disciplinary process", "People"),
    ("A.6.5", "Responsibilities after termination or change of employment", "People"),
    ("A.6.6", "Confidentiality or non-disclosure agreements", "People"),
    ("A.6.7", "Remote working", "People"),
    ("A.6.8", "Information security event reporting", "People"),
    # A.7 Physical (14)
    ("A.7.1", "Physical security perimeters", "Physical"),
    ("A.7.2", "Physical entry", "Physical"),
    ("A.7.3", "Securing offices, rooms and facilities", "Physical"),
    ("A.7.4", "Physical security monitoring", "Physical"),
    ("A.7.5", "Protecting against physical and environmental threats", "Physical"),
    ("A.7.6", "Working in secure areas", "Physical"),
    ("A.7.7", "Clear desk and clear screen", "Physical"),
    ("A.7.8", "Equipment siting and protection", "Physical"),
    ("A.7.9", "Security of assets off-premises", "Physical"),
    ("A.7.10", "Storage media", "Physical"),
    ("A.7.11", "Supporting utilities", "Physical"),
    ("A.7.12", "Cabling security", "Physical"),
    ("A.7.13", "Equipment maintenance", "Physical"),
    ("A.7.14", "Secure disposal or re-use of equipment", "Physical"),
    # A.8 Technological (34)
    ("A.8.1", "User endpoint devices", "Technological"),
    ("A.8.2", "Privileged access rights", "Technological"),
    ("A.8.3", "Information access restriction", "Technological"),
    ("A.8.4", "Access to source code", "Technological"),
    ("A.8.5", "Secure authentication", "Technological"),
    ("A.8.6", "Capacity management", "Technological"),
    ("A.8.7", "Protection against malware", "Technological"),
    ("A.8.8", "Management of technical vulnerabilities", "Technological"),
    ("A.8.9", "Configuration management", "Technological"),
    ("A.8.10", "Information deletion", "Technological"),
    ("A.8.11", "Data masking", "Technological"),
    ("A.8.12", "Data leakage prevention", "Technological"),
    ("A.8.13", "Information backup", "Technological"),
    ("A.8.14", "Redundancy of information processing facilities", "Technological"),
    ("A.8.15", "Logging", "Technological"),
    ("A.8.16", "Monitoring activities", "Technological"),
    ("A.8.17", "Clock synchronization", "Technological"),
    ("A.8.18", "Use of privileged utility programs", "Technological"),
    ("A.8.19", "Installation of software on operational systems", "Technological"),
    ("A.8.20", "Networks security", "Technological"),
    ("A.8.21", "Security of network services", "Technological"),
    ("A.8.22", "Segregation of networks", "Technological"),
    ("A.8.23", "Web filtering", "Technological"),
    ("A.8.24", "Use of cryptography", "Technological"),
    ("A.8.25", "Secure development life cycle", "Technological"),
    ("A.8.26", "Application security requirements", "Technological"),
    ("A.8.27", "Secure system architecture and engineering principles", "Technological"),
    ("A.8.28", "Secure coding", "Technological"),
    ("A.8.29", "Security testing in development and acceptance", "Technological"),
    ("A.8.30", "Outsourced development", "Technological"),
    ("A.8.31", "Separation of development, test and production environments", "Technological"),
    ("A.8.32", "Change management", "Technological"),
    ("A.8.33", "Test information", "Technological"),
    ("A.8.34", "Protection of information systems during audit testing", "Technological"),
]

EXCLUDED = {
    "A.7.1": "Excluded. Meridian operates no corporate premises: the company has been fully remote "
             "since founding and holds no leased office, data centre or warehouse. There is no "
             "perimeter to define. Production infrastructure sits in AWS, where physical perimeter "
             "security is the provider's responsibility under the shared responsibility model and is "
             "evidenced by AWS ISO 27001 and SOC 2 reports reviewed annually (see A.5.23).",
    "A.7.2": "Excluded. No company-controlled premises exist, so there are no entry points to "
             "control. AWS physical entry controls are inherited and evidenced as above.",
    "A.7.3": "Excluded. No offices, rooms or facilities are owned, leased or operated by Meridian.",
    "A.7.4": "Excluded. No company premises to monitor. Endpoint and cloud monitoring are covered "
             "by A.8.15 and A.8.16.",
    "A.7.6": "Excluded. No secure areas exist because no company premises exist. Access to "
             "production data is logical only and covered by A.8.2 and A.8.3.",
    "A.7.11": "Excluded. Meridian operates no facility requiring power, cooling or water supply. "
              "Utility resilience for production workloads is an AWS responsibility, evidenced in "
              "their availability commitments and reviewed under A.5.22.",
    "A.7.12": "Excluded. No company-owned cabling exists. Home-worker network security is addressed "
              "through A.6.7 remote working and A.8.1 endpoint controls.",
    "A.7.13": "Excluded. No company-operated equipment requires scheduled maintenance; laptops are "
              "leased with vendor replacement, and all production hardware is AWS-managed.",
    "A.8.30": "Excluded. All software development is performed by directly employed engineers. "
              "Meridian has no outsourced or contracted development relationship. This exclusion is "
              "reviewed at each management review; engaging a development contractor would make the "
              "control applicable immediately.",
}

# Control-specific justification for included controls; anything not listed uses a driver default.
JUSTIFY = {
    "A.5.1": "Risk R-01 (undefined security direction). Required by ISO 27001 Cl. 5.2 and customer DPAs.",
    "A.5.7": "Risk R-14 (targeted attacks on health data). Included at reduced scope: consumption of "
             "CISA and H-ISAC feeds rather than an in-house threat intelligence function, proportionate "
             "to a 60-person company.",
    "A.5.9": "Risk R-02 (unknown asset estate). Foundational to scope definition; prerequisite for A.5.12.",
    "A.5.12": "Risk R-03 (PHI handled at the same level as marketing content). Legal driver: HIPAA.",
    "A.5.13": "Risk R-03. Included but implemented as automated tagging in the data platform, not "
              "manual labelling of documents.",
    "A.5.15": "Risk R-04 (excessive standing access). Legal: HIPAA minimum necessary. Contractual: "
              "all enterprise customer agreements.",
    "A.5.18": "Risk R-04. Directly supported by the automated access review built in Project 3.",
    "A.5.19": "Risk R-08 (fourth-party exposure through subprocessors).",
    "A.5.23": "Risk R-09 (over-reliance on AWS with no verification). This control is where the "
              "inherited physical controls excluded in A.7 are actually verified.",
    "A.5.31": "Legal driver: HIPAA, state breach notification laws, GDPR for EU customers.",
    "A.5.34": "Legal driver: HIPAA and GDPR. Highest-consequence control in the ISMS.",
    "A.5.35": "Required for certification; delivered by an external assessor annually rather than an "
              "internal audit function Meridian does not have the headcount to staff.",
    "A.6.1": "Risk R-11 (insider access to PHI). Background screening proportionate to role sensitivity.",
    "A.6.7": "Risk R-12. Elevated importance: 100% of the workforce works remotely, so this control "
             "carries load that A.7 physical controls would carry elsewhere.",
    "A.7.5": "Included despite the remote model — covers environmental threats to the AWS regions in "
             "use, addressed through multi-AZ deployment and reviewed under A.8.14.",
    "A.7.7": "Included and applies in home working environments; enforced via screen-lock policy on "
             "managed endpoints.",
    "A.7.8": "Included at reduced scope: applies to company laptops in home offices, not facilities.",
    "A.7.9": "Included. Every company asset is permanently off-premises in this model, which makes "
             "this control more important here than in an office-based organisation.",
    "A.7.10": "Included at reduced scope: removable media is blocked by endpoint policy; the control "
              "documents that prohibition and the exception process.",
    "A.7.14": "Included. Laptop return and certified wipe on offboarding; AWS media destruction inherited.",
    "A.8.2": "Risk R-04 and R-05 (privileged standing access to production PHI). Highest-priority "
             "technical control; automated quarterly review implemented in Project 3.",
    "A.8.5": "Risk R-06. Phishing-resistant MFA on all production and administrative access.",
    "A.8.8": "Risk R-07 (unpatched dependencies). Weekly scanning with SLA by severity.",
    "A.8.11": "Risk R-10. Included: PHI is masked in all non-production environments — the control "
              "that makes A.8.33 workable.",
    "A.8.15": "Risk R-13. Legal: HIPAA audit controls. Also the evidence base for A.5.28.",
    "A.8.24": "Legal: HIPAA encryption addressable specification, treated as mandatory here. "
              "TLS 1.2+ in transit, AES-256 at rest.",
    "A.8.33": "Risk R-10. Production PHI is never used as test data; synthetic and masked datasets only.",
}

DEFAULT_DRIVERS = {
    "Organizational": "Risk assessment (ISMS risk register) and customer contractual requirements.",
    "People": "Risk assessment (insider and human-error risks); HIPAA workforce training requirement.",
    "Physical": "Risk assessment, scoped to remotely operated assets.",
    "Technological": "Risk assessment (ISMS risk register); HIPAA Security Rule technical safeguards.",
}

STATUS = {
    "A.5.7": "Partially implemented", "A.5.13": "Partially implemented",
    "A.5.28": "Planned Q1 2027", "A.5.30": "Partially implemented",
    "A.5.35": "Planned Q4 2026", "A.8.11": "Partially implemented",
    "A.8.12": "Planned Q1 2027", "A.8.16": "Partially implemented",
    "A.8.23": "Planned Q1 2027", "A.8.27": "Partially implemented",
    "A.8.29": "Partially implemented",
}

OWNERS = {"Organizational": "Head of Compliance", "People": "Head of People",
          "Physical": "Head of IT", "Technological": "CTO"}

wb = Workbook()

# ---------- Sheet 1: SoA ----------
ws = wb.active
ws.title = "Statement of Applicability"

headers = ["Control ID", "Control name", "Theme", "Applicable", "Justification",
           "Implementation status", "Control owner", "Evidence reference"]
ws.append(headers)

hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(color="FFFFFF", bold=True, size=11)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(vertical="center", horizontal="left")
ws.freeze_panes = "A2"

excl_fill = PatternFill("solid", fgColor="FCE4E4")
part_fill = PatternFill("solid", fgColor="FFF3CD")

for cid, name, theme in CONTROLS:
    if cid in EXCLUDED:
        applicable, justification, status, evidence = "No", EXCLUDED[cid], "N/A — excluded", "SoA rationale, Management Review MR-2026-03"
    else:
        applicable = "Yes"
        justification = JUSTIFY.get(cid, DEFAULT_DRIVERS[theme])
        status = STATUS.get(cid, "Implemented")
        evidence = "ISMS-EV-" + cid.replace("A.", "").replace(".", "")
    ws.append([cid, name, theme, applicable, justification, status, OWNERS[theme], evidence])
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c == 5))
    if applicable == "No":
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = excl_fill
    elif status != "Implemented":
        ws.cell(row=r, column=6).fill = part_fill

widths = [11, 46, 15, 11, 78, 22, 20, 22]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.auto_filter.ref = f"A1:H{ws.max_row}"

# ---------- Sheet 2: Summary ----------
ws2 = wb.create_sheet("Summary")
total = len(CONTROLS)
excluded = len(EXCLUDED)
included = total - excluded
partial = len([c for c in STATUS if c not in EXCLUDED])

rows = [
    ("ISO/IEC 27001:2022 Statement of Applicability", ""),
    ("Organisation", "Meridian Health Data, Inc. (scenario)"),
    ("ISMS scope", "The SaaS platform processing customer PHI, its supporting AWS production "
                   "environment, and the remote workforce operating it"),
    ("Version / date", "v1.0 — September 2026"),
    ("", ""),
    ("Total Annex A controls", total),
    ("Applicable", included),
    ("Excluded, with documented rationale", excluded),
    ("Applicable and fully implemented", included - partial),
    ("Applicable, partially implemented or planned", partial),
    ("", ""),
    ("Why exclusions matter", "A Statement of Applicability that marks all 93 controls applicable is "
                              "not a risk-driven SoA — it is a checklist. Each exclusion here is tied "
                              "to a factual attribute of the organisation, is reviewed at management "
                              "review, and becomes applicable the moment that attribute changes."),
    ("Excluded controls", ", ".join(sorted(EXCLUDED.keys(), key=lambda x: (len(x), x)))),
]
for row in rows:
    ws2.append(list(row))
ws2["A1"].font = Font(bold=True, size=14, color="1F3864")
for r in range(2, ws2.max_row + 1):
    ws2.cell(row=r, column=1).font = Font(bold=True)
    ws2.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 44
ws2.column_dimensions["B"].width = 88

wb.save("ISO27001-2022-Statement-of-Applicability.xlsx")
print(f"total={total} included={included} excluded={excluded}")
